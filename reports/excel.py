"""Save reports as Excel (.xlsx) files with openpyxl.

Amounts are written as numbers (Decimal, never float) with a money format, so
the owner can add them up in Excel.
"""
from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from services.errors import ValidationError
from services.report_service import DailySummary, summary_lines

_BOLD = Font(bold=True)
_TITLE = Font(bold=True, size=14)


def save_daily_summary(summary: DailySummary, path: str | Path, store_name: str,
                       currency: str) -> Path:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    money_fmt = f'"{currency}"#,##0.00;-"{currency}"#,##0.00'
    wb = Workbook()
    ws = wb.active
    ws.title = "Daily summary"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18

    ws.append([store_name])
    ws.cell(1, 1).font = _TITLE
    ws.append([f"Daily summary — {summary.day.strftime('%A, %d %B %Y')}"])
    for line in summary_lines(summary):
        if line.is_heading:
            ws.append([])
            ws.append([line.label])
            ws.cell(ws.max_row, 1).font = _BOLD
            continue
        ws.append([line.label, line.value])
        r = ws.max_row
        if isinstance(line.value, Decimal):
            ws.cell(r, 2).number_format = money_fmt
        if line.bold:
            ws.cell(r, 1).font = _BOLD
            ws.cell(r, 2).font = _BOLD
    try:
        wb.save(path)
    except PermissionError:
        raise ValidationError("The file could not be saved. If it is open in Excel, close it and "
                              "try again, or choose another name.") from None
    except OSError as exc:
        raise ValidationError(f"The file could not be saved ({exc.strerror or exc}).") from None
    return path
